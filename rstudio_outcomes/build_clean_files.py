"""Simplify the two uploaded, source-checked workbooks without changing values.
Usage: python build_clean_files.py continuous.xlsx binary.xlsx
Requires openpyxl. Outputs files alongside this script.
"""
from pathlib import Path
import sys,json,hashlib,math
from openpyxl import load_workbook,Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.utils import get_column_letter
BASE=Path(__file__).resolve().parent
if len(sys.argv)!=3:raise SystemExit(__doc__)
paths=[Path(p) for p in sys.argv[1:]]
books=[load_workbook(p,data_only=True) for p in paths]
def records(ws):
 rows=list(ws.values);return [dict(zip(rows[0],r)) for r in rows[1:] if any(v is not None for v in r)]
TASKS=[
 (0,'6MWD','01_six_minute_walk_distance.xlsx','walk_distance'),
 (0,'NYHA','02_nyha_class.xlsx','nyha'),
 (0,'Quality_of_life','03_quality_of_life.xlsx','quality_of_life'),
 (0,'QRS','04_qrs_duration.xlsx','qrs'),
 (0,'LVEF','05_lvef.xlsx','lvef'),
 (0,'LVESV_mL','06_lvesv_ml.xlsx','lvesv'),
 (0,'LVESV_percent','07_lvesv_percent_change.xlsx','lvesv'),
 (0,'LVESVi_mL_m2','08_lvesv_index_ml_m2.xlsx','lvesv'),
 (0,'Pacing_threshold','09_pacing_threshold.xlsx','threshold'),
 (1,'Composite','10_composite_outcome.xlsx','composite'),
 (1,'HF_hospitalization','11_heart_failure_hospitalization.xlsx','hhf'),
 (1,'All_cause_mortality','12_all_cause_mortality.xlsx','mortality')]
STAT_FIELDS=['median_csp','q1_csp','q3_csp','median_biv','q1_biv','q3_biv','se_csp','se_biv','ci_low_csp','ci_high_csp','ci_low_biv','ci_high_biv']
NUMERIC=['n_csp','mean_csp','sd_csp','n_biv','mean_biv','sd_biv','event_csp','event_biv']+STAT_FIELDS
MANIFEST=[];ALLIDS=[]
def add_sheet(wb,title,rows,keys=None):
 if not rows:return
 keys=keys or list(dict.fromkeys(k for r in rows for k in r))
 ws=wb.create_sheet(title);ws.append(keys)
 for r in rows:ws.append([r.get(k) for k in keys])
 for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='174A63');c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=30;ws.freeze_panes='B2';ws.sheet_view.zoomScale=90
 for i,k in enumerate(keys,1):
  width=16
  if k=='study':width=26
  elif k in ['statistic','measurement','follow_up','scale','definition','cohort_id','analysis_use','status','source_pdf','source_location']:width=28
  elif k in ['comparison','required_action','source_summary','sample_size_basis','analysis_population','design','explanation']:width=70
  elif k=='row_id':width=40
  ws.column_dimensions[get_column_letter(i)].width=width
 for row in ws.iter_rows(min_row=2):
  for cell in row:
   cell.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):cell.number_format='0.########'
  ws.row_dimensions[row[0].row].height=45 if title=='data' else 80
 tab=Table(displayName='Table_'+title,ref=ws.dimensions)
 tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
 ws.add_table(tab)
def slim(r,binary=False):
 vals=['n_csp','event_csp','n_biv','event_biv'] if binary else ['n_csp','mean_csp','sd_csp','n_biv','mean_biv','sd_biv']
 out={'study':r['study']}
 out.update({k:r.get('source_'+k) for k in vals})
 if not binary:
  out.update(statistic=r.get('statistic'),measurement=r.get('measurement_type'))
 out['follow_up']=r.get('follow_up');out['unit']=r.get('unit')
 if r.get('outcome')=='quality_of_life':out['scale']=r.get('scale')
 if binary:out['definition']=r.get('scale') or r.get('outcome')
 out['cohort_id']=r.get('overlap_group') or r['study']
 out['analysis_use']=r.get('pooling_eligibility')
 out['has_inputs']=all(out[k] is not None for k in vals) and (binary or r.get('statistic')=='mean and SD')
 out['row_id']=r.get('record_id')
 return out
for b,sheet,file,outcome in TASKS:
 source=records(books[b][sheet]);binary=b==1
 clean=[slim(r,binary) for r in source]
 wb=Workbook();wb.remove(wb.active);wb.properties.title=sheet+' - CSP versus BiVP';wb.properties.description='Source-checked values, simplified for RStudio. Blank means missing, never zero.'
 add_sheet(wb,'data',clean)
 # Additional statistical forms remain separate, keyed to the data row, never treated as raw SDs.
 stats=[]
 for r in source:
  other={k:r.get('source_'+k) for k in STAT_FIELDS}
  if any(x is not None for x in other.values()):stats.append({'study':r['study'],'statistic':r.get('statistic'),'measurement':r.get('measurement_type'),'unit':r.get('unit'),**other,'row_id':r['record_id']})
 if stats:
  keys=['study','statistic','measurement','unit']+[k for k in STAT_FIELDS if any(r.get(k) is not None for r in stats)]+['row_id']
  add_sheet(wb,'other_statistics',stats,keys)
 notes=[]
 for r in source:
  note={'study':r['study'],'status':r['status']}
  for k in ['comparison','required_action','sample_size_basis','analysis_population','design','source_csp_modality','source_pdf','pdf_pages','source_location','source_summary','same_estimand_as_figure','higher_is_better','group_n_csp','group_n_biv']:note[k]=r.get(k)
  note['row_id']=r['record_id'];notes.append(note)
 add_sheet(wb,'review_notes',notes)
 # Preserve labelled alternatives; do not duplicate them in primary data.
 extras=[r for r in records(books[b]['Additional_source_data']) if r['outcome']==outcome]
 if outcome=='lvesv':extras=[r for r in extras if r.get('unit')==source[0]['unit']]
 if extras:
  extra_rows=[]
  for r in extras:
   x=slim(r,binary);x['variant']=r.get('variant');x['source_pdf']=r.get('source_pdf');x['pdf_pages']=r.get('pdf_pages');x['source_location']=r.get('source_location');x['explanation']=r.get('source_summary')
   for k in STAT_FIELDS:
    if r.get('source_'+k) is not None:x[k]=r['source_'+k]
   extra_rows.append(x)
  add_sheet(wb,'alternatives',extra_rows)
 if not binary:
  contrasts=[r for r in records(books[0]['Study_contrasts']) if r['outcome']==outcome and (outcome!='lvesv' or r['unit']==source[0]['unit'])]
  if contrasts:add_sheet(wb,'reported_contrasts',contrasts)
 add_sheet(wb,'read_me',[
  {'item':'Import in RStudio','explanation':f'readxl::read_excel("{file}", sheet = "data")'},
  {'item':'Arms','explanation':'csp = conduction system pacing; biv = biventricular pacing. n means total participants, not non-events.'},
  {'item':'Missing data','explanation':'Blank numeric cells are unavailable source values. Do not replace them with zero or with unverified forest-plot numbers.'},
  {'item':'has_inputs','explanation':'TRUE only means all numerical arm inputs are present as mean/SD or event counts. It does not establish study independence or statistical eligibility. Check analysis_use and review_notes.'},
  {'item':'Statistics','explanation':'Only source mean/SD appear in main numeric columns. Medians, IQRs, model SEs and CIs remain in other_statistics when available. No conversions were performed.'},
  {'item':'Shared cohorts','explanation':'LEVEL-AT and CONSYST-CRT share participants. Never include both as independent evidence in the same pooled analysis.'},
  {'item':'Alternatives','explanation':'Additional source timepoints/populations are separate from data. Choose one appropriate record per study; do not append all alternatives as independent studies.'},
  {'item':'Compatibility','explanation':'Check units, measurement type, timepoint, paired versus independent design, and clinical endpoint definition before pooling.'}])
 path=BASE/file;wb.save(path)
 # Round-trip proof: clean numeric cells equal uploaded workbook source cells exactly.
 check=records(load_workbook(path,data_only=True)['data'])
 assert len(check)==len(source)
 for src,dst in zip(source,check):
  assert src['study']==dst['study'] and src['record_id']==dst['row_id']
  for k in ['n_csp','event_csp','n_biv','event_biv'] if binary else ['n_csp','mean_csp','sd_csp','n_biv','mean_biv','sd_biv']:
   assert src.get('source_'+k)==dst.get(k),(file,src['study'],k)
  assert not any(c.data_type=='f' for row in load_workbook(path)['data'] for c in row)
 ALLIDS.extend(r['row_id'] for r in check)
 MANIFEST.append(dict(file=file,outcome=outcome,rows=len(source),has_complete_numeric_inputs=sum(r['has_inputs'] for r in clean),source_workbook=paths[b].name,source_sheet=sheet,sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
assert len(ALLIDS)==96 and len(set(ALLIDS))==96
(BASE/'manifest.json').write_text(json.dumps({'source_workbooks':[{'name':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in paths],'outputs':MANIFEST,'total_primary_rows':96},indent=2)+'\n')
print('Created12 separate workbooks; all96 primary rows retained; all512 numerical arm cells match the uploaded source workbooks.')
for r in MANIFEST:print(r['file'],r['rows'],'rows')
